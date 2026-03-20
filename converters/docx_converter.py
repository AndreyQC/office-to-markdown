import os
from docx import Document
from io import BytesIO


def convert_docx(file_bytes: bytes, images_dir: str) -> str:
    doc = Document(BytesIO(file_bytes))
    markdown = []

    os.makedirs(images_dir, exist_ok=True)
    image_counter = 1

    for element in doc.element.body:
        if element.tag.endswith("p"):
            for para in doc.paragraphs:
                if para._element == element:
                    style = para.style.name if para.style else ""
                    text = para.text.strip()

                    if not text:
                        continue

                    if "Heading 1" in style:
                        markdown.append(f"# {text}\n")
                    elif "Heading 2" in style:
                        markdown.append(f"## {text}\n")
                    elif "Heading 3" in style:
                        markdown.append(f"### {text}\n")
                    else:
                        markdown.append(f"{text}\n")
                    break

        elif element.tag.endswith("tbl"):
            for table in doc.tables:
                if table._tbl == element:
                    table_md = _convert_table(table)
                    markdown.append(table_md)
                    break

    for rel in doc.part.rels.values():
        target_ref = getattr(rel, "target_ref", "") or ""
        if "image" not in target_ref:
            continue

        # External relationships (linked web images) do not expose target_part.
        if getattr(rel, "is_external", False):
            continue

        target_part = getattr(rel, "target_part", None)
        if target_part is None:
            continue

        image_data = target_part.blob
        image_ext = target_ref.split(".")[-1]
        image_name = f"img_{image_counter}.{image_ext}"
        image_path = os.path.join(images_dir, image_name)

        with open(image_path, "wb") as f:
            f.write(image_data)

        markdown.append(f"\n![image](images/{image_name})\n")
        image_counter += 1

    return "\n".join(markdown)


def _convert_table(table) -> str:
    rows = []

    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)

    if not rows:
        return ""

    result = []
    result.append("| " + " | ".join(rows[0]) + " |")
    result.append("| " + " | ".join(["---"] * len(rows[0])) + " |")

    for row in rows[1:]:
        result.append("| " + " | ".join(row) + " |")

    return "\n".join(result) + "\n"
