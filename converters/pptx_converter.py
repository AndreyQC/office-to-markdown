import os
from pptx import Presentation
from io import BytesIO


def convert_pptx(file_bytes: bytes, images_dir: str) -> str:
    prs = Presentation(BytesIO(file_bytes))
    markdown = []

    os.makedirs(images_dir, exist_ok=True)
    image_counter = 1

    for slide_num, slide in enumerate(prs.slides, 1):
        markdown.append(f"## Слайд {slide_num}\n")

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                markdown.append(shape.text.strip())
                markdown.append("")

            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    image = shape.image
                    image_bytes = image.blob
                    image_ext = image.ext
                    image_name = f"img_{image_counter}.{image_ext}"
                    image_path = os.path.join(images_dir, image_name)

                    with open(image_path, "wb") as f:
                        f.write(image_bytes)

                    markdown.append(f"![image](images/{image_name})\n")
                    image_counter += 1
                except Exception:
                    pass

        markdown.append("---\n")

    return "\n".join(markdown)
